"""把 svglib 生成的 SVG 逐元素还原为原生、可编辑的 PPT 形状与文本框。"""
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525  # 1px = 9525 EMU（96 dpi）
FONT = 'Microsoft YaHei'
NS = '{http://www.w3.org/2000/svg}'


def emu(v):
    return Emu(int(round(float(v) * PX)))


def rgb(h):
    return RGBColor.from_string(h.lstrip('#').upper())


def apply_fill(shape, fill, opacity=1.0):
    if not fill or fill == 'none':
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if opacity < 1:
        clr = shape.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(opacity * 100000)))


def apply_line(shape, el):
    stroke = el.get('stroke')
    if not stroke or stroke == 'none':
        shape.line.fill.background()
        return
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = emu(el.get('stroke-width', '1'))
    if el.get('stroke-dasharray'):
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = shape.line._get_or_add_ln()
    if el.get('stroke-linejoin') == 'round':
        etree.SubElement(ln, qn('a:round'))
    for attr, tag in (('marker-start', 'a:headEnd'), ('marker-end', 'a:tailEnd')):
        if el.get(attr):
            e = etree.SubElement(ln, qn(tag))
            e.set('type', 'triangle')
            e.set('w', 'med')
            e.set('len', 'med')


def strip_style(shape):
    """去掉主题样式引用（默认阴影等），只保留显式设置的填充和线条。"""
    style = shape._element.find(qn('p:style'))
    if style is not None:
        shape._element.remove(style)


def points(s):
    return [tuple(float(v) for v in p.split(',')) for p in s.split()]


def freeform(slide, pts, close):
    fb = slide.shapes.build_freeform(emu(pts[0][0]), emu(pts[0][1]), scale=1.0)
    fb.add_line_segments([(emu(x), emu(y)) for x, y in pts[1:]], close=close)
    return fb.convert_to_shape()


def set_run_font(run, size, color, bold):
    f = run.font
    f.size = Pt(size * 0.75)
    f.bold = bool(bold)
    f.color.rgb = rgb(color)
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rpr, qn('a:ea'))
    ea.set('typeface', FONT)


def add_text(slide, meta):
    x, y, w, h = meta['box']
    tb = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[meta['valign']]
    size = meta['size']
    align = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[meta['align']]
    for i, runs in enumerate(meta['paras']):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = Pt(size * meta['lh'] * 0.75)
        if i and meta['gap']:
            p.space_before = Pt(meta['gap'] * 0.75)
        if meta.get('bullet'):
            ind = int(size * 1.1 * PX)
            ppr = p._p.get_or_add_pPr()
            ppr.set('marL', str(ind))
            ppr.set('indent', str(-ind))
            bc = etree.SubElement(ppr, qn('a:buClr'))
            etree.SubElement(bc, qn('a:srgbClr')).set('val', meta['bullet_color'].lstrip('#').upper())
            etree.SubElement(ppr, qn('a:buFont')).set('typeface', 'Arial')
            etree.SubElement(ppr, qn('a:buChar')).set('char', meta['bullet'])
        for text, color, bold in runs:
            r = p.add_run()
            r.text = text
            set_run_font(r, size, color, bold)
    return tb


def convert_slide(slide, svg_path):
    root = etree.parse(str(svg_path)).getroot()
    for el in root:
        tag = el.tag.replace(NS, '')
        if tag == 'rect':
            x, y, w, h = (float(el.get(k)) for k in ('x', 'y', 'width', 'height'))
            rx = float(el.get('rx', 0))
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if rx else MSO_SHAPE.RECTANGLE
            sh = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
            if rx:
                sh.adjustments[0] = min(0.5, rx / min(w, h))
        elif tag == 'circle':
            cx, cy, r = (float(el.get(k)) for k in ('cx', 'cy', 'r'))
            sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(cx - r), emu(cy - r), emu(2 * r), emu(2 * r))
        elif tag == 'polygon':
            sh = freeform(slide, points(el.get('points')), close=True)
        elif tag == 'polyline':
            pts = points(el.get('points'))
            if len(pts) == 2:
                (x1, y1), (x2, y2) = pts
                sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
            else:
                sh = freeform(slide, pts, close=False)
                sh.fill.background()
            apply_line(sh, el)
            strip_style(sh)
            continue
        elif tag == 'text':
            add_text(slide, json.loads(el.get('data-pptx')))
            continue
        else:
            continue
        apply_fill(sh, el.get('fill'), float(el.get('fill-opacity', 1)))
        apply_line(sh, el)
        strip_style(sh)
        if sh.has_text_frame:
            sh.text_frame.text = ''


def build(svg_files, out):
    prs = Presentation()
    prs.slide_width = Emu(1280 * PX)
    prs.slide_height = Emu(720 * PX)
    blank = prs.slide_layouts[6]
    for f in svg_files:
        convert_slide(prs.slides.add_slide(blank), f)
    prs.save(out)


if __name__ == '__main__':
    files = sorted(Path(sys.argv[1]).glob('*.svg'))
    build(files, sys.argv[2])
    print(f'{len(files)} slides -> {sys.argv[2]}')
